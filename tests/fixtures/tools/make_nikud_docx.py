"""Build nikud-quotes.docx (dev-only fixture) with python-docx.

Content exercises the .docx extractor and the normalizer: Hebrew with
nikud (points), gershayim/geresh (U+05F4/U+05F3), ASCII and curly quotes,
a maqaf, a non-breaking hyphen, a tab, a manual line break inside a
paragraph, right-to-left paragraphs and a 2x2 table.

Usage: python3 make_nikud_docx.py <out.docx>
"""

import sys

from docx import Document
from docx.oxml import OxmlElement

PARAGRAPHS = [
    "בְּרֵאשִׁית בָּרָא אֱלֹהִים אֵת הַשָּׁמַיִם וְאֵת הָאָרֶץ",
    "עו״ד ישראל ישראלי ג׳ורג׳ בע״מ",
    "הוא אמר: “שלום” ו‘להתראות’ – \"בסדר\" ו'כן'",
    "תל־אביב־יפו",
]


def set_rtl(paragraph):
    """Mark a paragraph as right-to-left (w:bidi)."""
    p_pr = paragraph._p.get_or_add_pPr()
    bidi = OxmlElement("w:bidi")
    p_pr.append(bidi)


def add_no_break_hyphen(run):
    """Append a w:noBreakHyphen to a run."""
    run._r.append(OxmlElement("w:noBreakHyphen"))


def main():
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    document = Document()
    document.core_properties.author = "fixture"

    for text in PARAGRAPHS:
        set_rtl(document.add_paragraph(text))

    # Manual line break, tab and a non-breaking hyphen in one paragraph.
    paragraph = document.add_paragraph()
    set_rtl(paragraph)
    run = paragraph.add_run("שורה ראשונה")
    run.add_break()
    paragraph.add_run("שורה שנייה")
    paragraph.add_run().add_tab()
    run = paragraph.add_run("57879")
    add_no_break_hyphen(run)
    paragraph.add_run("06-22")

    table = document.add_table(rows=2, cols=2)
    cells = [["תא א1", "תא א2"], ["תא ב1", "תא ב2"]]
    for r, row in enumerate(cells):
        for c, value in enumerate(row):
            table.cell(r, c).text = value

    set_rtl(document.add_paragraph("סוף המסמך"))
    document.save(sys.argv[1])


if __name__ == "__main__":
    main()
